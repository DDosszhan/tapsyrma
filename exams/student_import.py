from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def read_student_emails(upload):
    if not upload.name.lower().endswith(".xlsx") or upload.size > 2_000_000:
        raise ValidationError("Upload an .xlsx file up to 2 MB.")
    try:
        book = load_workbook(upload, read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, EOFError, ParseError) as exc:
        raise ValidationError("Could not open the Excel file.") from exc
    try:
        try:
            rows, seen = [], set()
            for line, cells in enumerate(book.active.iter_rows(values_only=True), start=1):
                if line > 1001:
                    raise ValidationError("A file may contain up to 1,000 students.")
                raw = cells[0] if cells else None
                if line == 1 and str(raw).strip().lower() in {"email", "e-mail", "почта"}:
                    continue
                if raw is None or not str(raw).strip():
                    if any(value is not None and str(value).strip() for value in cells[1:]):
                        raise ValidationError(f"Row {line}: enter an email address in the first column.")
                    continue
                email = str(raw).strip().lower()
                try:
                    validate_email(email)
                except ValidationError as exc:
                    raise ValidationError(f"Row {line}: invalid email address.") from exc
                if len(email) > 150:
                    raise ValidationError(f"Row {line}: email address is too long (maximum 150 characters).")
                if email in seen:
                    raise ValidationError(f"Row {line}: duplicate email address in this file.")
                seen.add(email)
                name = str(cells[1]).strip() if len(cells) > 1 and cells[1] is not None else email
                if not name or len(name) > 200:
                    raise ValidationError(f"Row {line}: name must be at most 200 characters.")
                rows.append((email, name))
                if len(rows) > 1000:
                    raise ValidationError("A file may contain up to 1,000 students.")
            if not rows:
                raise ValidationError("The file contains no student email addresses.")
            return rows
        except (BadZipFile, OSError, ValueError, KeyError, EOFError, IndexError, ParseError) as exc:
            raise ValidationError("Could not read the Excel rows.") from exc
    finally:
        book.close()
